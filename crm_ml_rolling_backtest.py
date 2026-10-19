#!/usr/bin/env python3
"""Rolling-origin evaluator for crm_ml_bayes_uq_reviewed.py.

Designed to keep the *runner* lightweight: it uses only the standard library and
openpyxl. It deliberately does not import NumPy, pandas, JAX, or NumPyro in the
parent process. The model script is launched as a child process per fold, and
its combined stdout/stderr is streamed to the terminal and to that fold's run.log.

Each fold trains only through `origin` and evaluates the next `horizon` rows.
Origins advance by `step` rows. All candidate fits use the same fold boundaries.
"""
from __future__ import annotations

import argparse
import csv
import json
import math
import statistics
import subprocess
import sys
import traceback
from pathlib import Path

print("[startup] rolling-origin runner entered", flush=True)
try:
    from openpyxl import Workbook, load_workbook
except Exception:
    print("[startup] failed to import openpyxl", file=sys.stderr, flush=True)
    traceback.print_exc()
    raise
print("[startup] lightweight dependencies loaded (openpyxl); numerical model imports are deferred to each child process", flush=True)


class WorkbookData:
    def __init__(self, time, production, injection, producers, injectors):
        self.time = time
        self.production = production
        self.injection = injection
        self.producers = producers
        self.injectors = injectors

    @property
    def n_prod(self):
        return len(self.producers)

    @property
    def n_inj(self):
        return len(self.injectors)


def _find_sheet(workbook, aliases):
    by_lower = {name.strip().lower(): name for name in workbook.sheetnames}
    for alias in aliases:
        if alias.lower() in by_lower:
            return workbook[by_lower[alias.lower()]]
    raise ValueError(f"Could not find a worksheet matching {aliases}; available sheets: {workbook.sheetnames}")


def _numeric_sheet(sheet):
    """Match the source loader's format: first row labels, then fully numeric rows.

    Unit rows (e.g. '[days]' and '[MSTB/day]') are ignored. A partially blank or
    non-numeric data row is skipped, just as the source model's _numeric_block does.
    """
    rows = list(sheet.iter_rows(values_only=True))
    rows = [list(row) for row in rows if row and any(v is not None for v in row)]
    if len(rows) < 2:
        raise ValueError(f"Worksheet {sheet.title!r} does not contain headers plus numeric data")
    labels = [str(v).strip() if v is not None else "" for v in rows[0]]
    numeric = []
    for row in rows[1:]:
        if len(row) < len(labels):
            row = row + [None] * (len(labels) - len(row))
        converted = []
        ok = True
        for value in row[:len(labels)]:
            if value is None or isinstance(value, bool):
                ok = False
                break
            try:
                number = float(value)
            except (TypeError, ValueError):
                ok = False
                break
            if not math.isfinite(number):
                ok = False
                break
            converted.append(number)
        if ok:
            numeric.append(converted)
    if not numeric:
        raise ValueError(f"Worksheet {sheet.title!r} has no fully numeric data rows")
    return labels, numeric


def load_workbook_data(path: Path) -> WorkbookData:
    print(f"[preflight] reading workbook without importing NumPy/pandas: {path}", flush=True)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        _, trows = _numeric_sheet(_find_sheet(wb, ["time", "t"]))
        producers, production = _numeric_sheet(_find_sheet(wb, ["q_observed", "production"]))
        injectors, injection = _numeric_sheet(_find_sheet(wb, ["w_injection", "injection"]))
    finally:
        wb.close()
    time = [row[0] for row in trows]
    if not (len(time) == len(production) == len(injection)):
        raise ValueError(
            f"Workbook series lengths differ: time={len(time)}, production={len(production)}, injection={len(injection)}"
        )
    if any(b <= a for a, b in zip(time, time[1:])):
        raise ValueError("Time values must be strictly increasing")
    if any(x < 0 for row in production for x in row) or any(x < 0 for row in injection for x in row):
        raise ValueError("Production and injection values must be non-negative")
    if any(len(row) != len(producers) for row in production):
        raise ValueError("Production rows do not match the producer header width")
    if any(len(row) != len(injectors) for row in injection):
        raise ValueError("Injection rows do not match the injector header width")
    return WorkbookData(time, production, injection, producers, injectors)


def write_truncated_workbook(data: WorkbookData, end: int, path: Path) -> None:
    """Write rows [0:end] to the three canonical sheets used by the model script."""
    wb = Workbook(write_only=True)
    ws_time = wb.create_sheet("time")
    ws_time.append(["time"])
    for value in data.time[:end]:
        ws_time.append([value])
    ws_prod = wb.create_sheet("q_observed")
    ws_prod.append(data.producers)
    for row in data.production[:end]:
        ws_prod.append(row)
    ws_inj = wb.create_sheet("w_injection")
    ws_inj.append(data.injectors)
    for row in data.injection[:end]:
        ws_inj.append(row)
    wb.save(path)


def make_parser():
    p = argparse.ArgumentParser(description="Rolling-origin evaluation of implemented CRM-ML models")
    p.add_argument("--data", required=True, help="Original CRM workbook")
    p.add_argument("--model-script", default=str(Path(__file__).with_name("crm_ml_bayes_uq_reviewed.py")),
                   help="Path to the CRM-ML Python script")
    p.add_argument("--out", default="rolling_origin_results", help="Output directory")
    p.add_argument("--horizon", type=int, default=12, help="Forecast horizon in rows (12 if monthly)")
    p.add_argument("--step", type=int, default=6, help="Advance between forecast origins in rows")
    p.add_argument("--initial-history", type=int, default=60, help="Minimum training-history length in rows")
    p.add_argument("--max-folds", type=int, default=4, help="Latest N valid origins; 0 runs all valid origins")
    p.add_argument("--ml", nargs="+", default=["xgb"], choices=["xgb", "elm", "mlp", "nusvm"])
    p.add_argument("--likelihood", default="normal", choices=["normal", "ar1", "lognormal"])
    p.add_argument("--tau-min", type=float, default=1 / 30.44)
    p.add_argument("--starts", type=int, default=8)
    p.add_argument("--max-iter", type=int, default=1000)
    p.add_argument("--chains", type=int, default=4)
    p.add_argument("--warmup", type=int, default=1000)
    p.add_argument("--samples", type=int, default=1000)
    p.add_argument("--tau-prior-sd", type=float, default=2.0)
    p.add_argument("--alpha", type=float, default=1.0)
    p.add_argument("--p-convention", default="percentile", choices=["percentile", "reserves"])
    p.add_argument("--ml-draws", type=int, default=1000)
    p.add_argument("--seed", type=int, default=2025)
    p.add_argument("--stop-on-error", action="store_true")
    return p


def read_csv_rows(path: Path):
    with path.open("r", newline="", encoding="utf-8-sig") as fh:
        return list(csv.DictReader(fh))


def write_csv_rows(path: Path, rows):
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    # Put fields from all rows in a stable first-seen order.
    fields = []
    seen = set()
    for row in rows:
        for key in row:
            if key not in seen:
                seen.add(key)
                fields.append(key)
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def maybe_float(value):
    try:
        x = float(value)
        return x if math.isfinite(x) else None
    except (TypeError, ValueError):
        return None


def summarise_rows(rows, output_path: Path):
    good = [r for r in rows if r.get("status") == "OK" and r.get("model")]
    if not good:
        print("[summary] no successful model metrics to aggregate", flush=True)
        return
    metric_names = ["MAE", "RMSE", "R2_det", "R2_corr", "PICP_80", "width_80", "IS_80",
                    "PICP_95", "width_95", "IS_95", "CRPS"]
    grouped = {}
    for r in good:
        grouped.setdefault(r["model"], []).append(r)
    summary = []
    for model, model_rows in grouped.items():
        record = {"model": model, "n_folds": len(model_rows)}
        for metric in metric_names:
            vals = [v for v in (maybe_float(r.get(metric)) for r in model_rows) if v is not None]
            if vals:
                record[f"{metric}_mean"] = statistics.fmean(vals)
                record[f"{metric}_std"] = statistics.stdev(vals) if len(vals) > 1 else 0.0
                record[f"{metric}_count"] = len(vals)
        summary.append(record)
    write_csv_rows(output_path, summary)


def run_child(cmd, log_path: Path):
    """Run one model fold and stream combined output to console and log."""
    with log_path.open("w", encoding="utf-8", errors="replace") as log_file:
        log_file.write("COMMAND: " + subprocess.list2cmdline(cmd) + "\n\n")
        log_file.flush()
        try:
            proc = subprocess.Popen(
                cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                text=True, encoding="utf-8", errors="replace", bufsize=1,
            )
        except Exception:
            traceback.print_exc(file=log_file)
            traceback.print_exc()
            return 9009
        assert proc.stdout is not None
        for line in proc.stdout:
            print("    " + line, end="", flush=True)
            log_file.write(line)
            log_file.flush()
        code = proc.wait()
        print(f"    [child exit code: {code}]", flush=True)
        log_file.write(f"\n[child exit code: {code}]\n")
        return code


def main():
    args = make_parser().parse_args()
    data_path = Path(args.data).expanduser().resolve()
    model_path = Path(args.model_script).expanduser().resolve()
    out = Path(args.out).expanduser().resolve()
    print(f"[preflight] runner Python: {sys.executable}", flush=True)
    print(f"[preflight] data exists: {data_path.is_file()} ({data_path})", flush=True)
    print(f"[preflight] model exists: {model_path.is_file()} ({model_path})", flush=True)
    if not data_path.is_file():
        raise FileNotFoundError(f"Data workbook not found: {data_path}")
    if not model_path.is_file():
        raise FileNotFoundError(f"Model script not found: {model_path}")
    if args.horizon < 1 or args.step < 1:
        raise ValueError("--horizon and --step must both be positive")
    if args.initial_history < 10:
        raise ValueError("--initial-history must be at least 10 rows")
    if args.max_folds < 0:
        raise ValueError("--max-folds must be 0 or positive")

    data = load_workbook_data(data_path)
    n = len(data.time)
    print(f"[preflight] workbook loaded: {n} rows, {data.n_prod} producers, {data.n_inj} injectors", flush=True)
    origins = list(range(args.initial_history, n - args.horizon + 1, args.step))
    if args.max_folds > 0:
        origins = origins[-args.max_folds:]
    if len(origins) < 2:
        raise ValueError(
            f"Only {len(origins)} eligible origin(s) for N={n}, initial history={args.initial_history}, "
            f"horizon={args.horizon}; need at least two. Adjust --initial-history, --horizon or --step."
        )

    out.mkdir(parents=True, exist_ok=True)
    settings = vars(args).copy()
    settings.update({"n_rows": n, "n_producers": data.n_prod, "n_injectors": data.n_inj,
                     "origins_rows": origins, "model_script_resolved": str(model_path),
                     "data_resolved": str(data_path), "runner_python": sys.executable})
    (out / "backtest_settings.json").write_text(json.dumps(settings, indent=2), encoding="utf-8")
    print(f"[preflight] eligible origins: {origins}; output directory: {out}", flush=True)

    all_rows = []
    for fold, origin in enumerate(origins, start=1):
        end = origin + args.horizon
        fold_dir = out / f"fold_{fold:02d}_origin_{origin:04d}"
        fold_dir.mkdir(parents=True, exist_ok=True)
        input_path = fold_dir / "fold_input.xlsx"
        print(f"[{fold}/{len(origins)}] writing fold workbook for rows 0:{end}", flush=True)
        write_truncated_workbook(data, end, input_path)
        model_out = fold_dir / "model_output"
        cmd = [sys.executable, "-u", str(model_path),
               "--data", str(input_path), "--forecast", str(args.horizon), "--out", str(model_out),
               "--tau-min", str(args.tau_min), "--starts", str(args.starts),
               "--max-iter", str(args.max_iter), "--likelihood", args.likelihood,
               "--chains", str(args.chains), "--warmup", str(args.warmup),
               "--samples", str(args.samples), "--tau-prior-sd", str(args.tau_prior_sd),
               "--alpha", str(args.alpha), "--p-convention", args.p_convention,
               "--seed", str(args.seed + fold - 1), "--ml-draws", str(args.ml_draws),
               "--ml", *list(dict.fromkeys(args.ml))]
        log_path = fold_dir / "run.log"
        print(f"[{fold}/{len(origins)}] origin row {origin}; forecast rows {origin}:{end}; launching child model", flush=True)
        return_code = run_child(cmd, log_path)
        if return_code != 0:
            msg = f"Fold {fold} failed (exit={return_code}); inspect {log_path}"
            print("[error] " + msg, file=sys.stderr, flush=True)
            all_rows.append({"fold": fold, "origin_row": origin, "end_row_exclusive": end,
                             "status": "FAILED", "return_code": return_code, "log_path": str(log_path)})
            if args.stop_on_error:
                break
            continue

        metric_path = model_out / "metrics.csv"
        if not metric_path.is_file():
            msg = f"Fold {fold} exited successfully but did not create metrics.csv; inspect {log_path}"
            print("[error] " + msg, file=sys.stderr, flush=True)
            all_rows.append({"fold": fold, "origin_row": origin, "end_row_exclusive": end,
                             "status": "NO_METRICS", "return_code": 0, "log_path": str(log_path)})
            if args.stop_on_error:
                break
            continue

        holdout_rows = [r for r in read_csv_rows(metric_path)
                        if r.get("period") == "holdout" and r.get("producer") == "FIELD(pooled)"]
        if not holdout_rows:
            print(f"[warning] no pooled holdout rows found in fold {fold}; inspect {metric_path}", flush=True)
        for metric in holdout_rows:
            metric.update({"fold": fold, "origin_row": origin,
                           "history_end_time": data.time[origin - 1],
                           "forecast_start_time": data.time[origin],
                           "forecast_end_time": data.time[end - 1], "status": "OK"})
            all_rows.append(metric)
        print(f"[{fold}/{len(origins)}] saved {len(holdout_rows)} pooled holdout metric rows", flush=True)

    fold_file = out / "rolling_origin_fold_metrics.csv"
    write_csv_rows(fold_file, all_rows)
    summarise_rows(all_rows, out / "rolling_origin_summary.csv")
    print(f"[finished] fold metrics: {fold_file}", flush=True)
    print(f"[finished] outputs: {out}", flush=True)
    failed = sum(1 for r in all_rows if r.get("status") in {"FAILED", "NO_METRICS"})
    if failed:
        print(f"[finished] warning: {failed} fold(s) failed; see recorded log_path values", flush=True)
        return 1
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        print("\n[interrupted] KeyboardInterrupt received; completed folds and logs were preserved.", file=sys.stderr, flush=True)
        raise SystemExit(130)
    except SystemExit:
        raise
    except BaseException:
        print("[fatal] Rolling-origin runner crashed with an exception:", file=sys.stderr, flush=True)
        traceback.print_exc()
        raise SystemExit(1)
